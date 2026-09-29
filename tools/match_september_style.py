# -*- coding: utf-8 -*-
"""10월 이후 갱신보험료 파일을 2026.09 파일과 같은 서식으로 맞춘다.

2026.09 파일의 서식을 기준으로 삼는다.
  - 제목행: Calibri Light 11 (굵게 아님) + 살구색 배경(F8CBAD) + 가운데
  - 본문:   Calibri 11 + 가로·세로 가운데
  - 전 셀 얇은 테두리
  - 금액 열만 오른쪽 정렬 (10·11월 방식 유지)
  - 금액 열 숫자 서식 #,##0.00
  - 인쇄 방향 가로

9월과 달라지는 점은 금액 열 우측 정렬 하나뿐이다.
값과 수식은 건드리지 않는다.

사용법:
    python match_september_style.py "2026.11.xlsx"
    python match_september_style.py "C:\\...\\갱신보험료 납입" --backup
    python match_september_style.py . --reference "2026.09.xlsx"   # 기준 파일에서 폭을 직접 읽기
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl 이 필요합니다.  먼저 실행:  pip install openpyxl")


# ------------------------------------------------- 2026.09 파일에서 읽어낸 기준값
HEADER_FONT = Font(name="Calibri Light", size=11, bold=False)
HEADER_FILL = PatternFill("solid", fgColor="FFF8CBAD")
BODY_FONT = Font(name="Calibri", size=11)

CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

_LINE = Side(style="thin")
BORDER = Border(left=_LINE, right=_LINE, top=_LINE, bottom=_LINE)

# 제목 이름 -> 열 너비. 열 순서가 달라져도 제목으로 찾으므로 안전하다.
# 9월 파일에 폭이 지정돼 있던 열만 들어 있다. 나머지는 엑셀 기본값 그대로 둔다.
REFERENCE_WIDTHS = {
    "provider": 9.0,
    "policy no.": 12.875,
    "n차년": 9.0,
    "owner": 19.5,
    "premium + levy (usd)": 15.75,
    "paid to date": 12.875,
    "fpd (usd)": 9.0,
    "송금일": 15.625,
}

AMOUNT_FORMAT = "#,##0.00"

# 제목에 이 말이 들어간 열은 금액 열로 보고 우측 정렬 + #,##0.00
_MONEY_HEADER = (
    "금액", "보험료", "납입액", "합계", "잔액", "대출", "환산",
    "premium", "levy", "loan", "amount", "balance", "total", "fee",
)

# 제목에 이 말이 들어간 열은 숫자여도 식별번호로 보고 가운데
_ID_HEADER = (
    "번호", "증권", "코드", "회차", "차년", "연령", "나이", "건수",
    "no.", "policy", "code", "phone", "count", "age",
)


def _norm(value) -> str:
    return str(value or "").strip().lower()


def _money_columns(ws) -> set:
    """금액 열 번호 집합. 제목으로 판정하고, 본문이 숫자인지로 한 번 더 거른다."""
    cols = set()
    for col in range(1, ws.max_column + 1):
        title = _norm(ws.cell(row=1, column=col).value)
        if not title or any(k in title for k in _ID_HEADER):
            continue
        if not any(k in title for k in _MONEY_HEADER):
            continue

        filled = numeric = 0
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=col).value
            if v is None or v == "":
                continue
            filled += 1
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                numeric += 1
        # 값이 아직 없는 열(예: 다음 달 예정 금액)도 금액 열로 인정한다
        if filled == 0 or numeric * 2 >= filled:
            cols.add(col)
    return cols


def _widths_from_reference(path: Path) -> dict:
    """기준 파일(2026.09)에서 제목 이름별 열 너비를 직접 읽어온다."""
    ws = load_workbook(path).worksheets[0]
    widths = {}
    for col in range(1, ws.max_column + 1):
        title = _norm(ws.cell(row=1, column=col).value)
        dim = ws.column_dimensions.get(get_column_letter(col))
        if title and dim is not None and dim.width:
            widths[title] = dim.width
    return widths


def format_sheet(ws, widths: dict) -> None:
    if ws.max_row < 1 or ws.sheet_state != "visible":
        return

    money_cols = _money_columns(ws)

    merged = {
        (r, c)
        for rng in ws.merged_cells.ranges
        for r in range(rng.min_row, rng.max_row + 1)
        for c in range(rng.min_col, rng.max_col + 1)
    } - {(rng.min_row, rng.min_col) for rng in ws.merged_cells.ranges}

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if (cell.row, cell.column) in merged:
                continue

            if cell.row == 1:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = CENTER
            else:
                cell.font = BODY_FONT
                if cell.column in money_cols:
                    cell.alignment = RIGHT
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = AMOUNT_FORMAT
                else:
                    cell.alignment = CENTER

            cell.border = BORDER

    # 열 너비: 제목 이름으로 기준값을 찾아 적용. 기준에 없는 열은 손대지 않는다.
    for col in range(1, ws.max_column + 1):
        title = _norm(ws.cell(row=1, column=col).value)
        if title in widths:
            ws.column_dimensions[get_column_letter(col)].width = widths[title]

    ws.page_setup.orientation = "landscape"


def format_file(path: Path, widths: dict, backup: bool) -> str:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        format_sheet(ws, widths)
    if backup:
        shutil.copy2(path, path.with_suffix(path.suffix + ".bak"))
    wb.save(path)
    return f"{path.name}  ({len(wb.worksheets)}개 시트)"


def main() -> int:
    ap = argparse.ArgumentParser(description="갱신보험료 파일을 2026.09 서식으로 통일")
    ap.add_argument("target", nargs="?", default=".", help="대상 폴더 또는 .xlsx 파일")
    ap.add_argument("--reference", help="기준 파일 경로 (생략하면 내장된 9월 기준값 사용)")
    ap.add_argument("--recursive", action="store_true", help="하위 폴더까지")
    ap.add_argument("--backup", action="store_true", help="원본을 .xlsx.bak 으로 백업")
    args = ap.parse_args()

    widths = dict(REFERENCE_WIDTHS)
    if args.reference:
        ref = Path(args.reference).expanduser()
        if not ref.is_file():
            sys.exit(f"기준 파일을 찾을 수 없습니다: {ref}")
        widths = _widths_from_reference(ref)
        print(f"기준 파일에서 열 너비 {len(widths)}개를 읽었습니다: {ref.name}")

    target = Path(args.target).expanduser()
    if not target.exists():
        sys.exit(f"경로를 찾을 수 없습니다: {target}")

    if target.is_file():
        files = [target]
    else:
        pattern = "**/*.xlsx" if args.recursive else "*.xlsx"
        files = sorted(p for p in target.glob(pattern) if not p.name.startswith("~$"))
        if args.reference:
            ref_resolved = Path(args.reference).expanduser().resolve()
            files = [f for f in files if f.resolve() != ref_resolved]

    if not files:
        sys.exit(f"처리할 .xlsx 파일이 없습니다: {target}")

    ok, failed = 0, []
    for f in files:
        try:
            print("  [완료]", format_file(f, widths, args.backup))
            ok += 1
        except Exception as exc:
            failed.append((f.name, exc))
            print("  [실패]", f.name, "-", exc)

    print(f"\n총 {len(files)}개 중 {ok}개 처리 완료.")
    if failed:
        print("실패 목록:")
        for name, exc in failed:
            print(f"  - {name}: {exc}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
